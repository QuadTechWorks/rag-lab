from __future__ import annotations
from core.registry import LOADERS
from core.interfaces.base_loader import BaseLoader
from core.models.documents import LoadedDocument


@LOADERS.register(provider="langchain", name="pptx")
class LangChainPPTXLoader(BaseLoader):
    """Load PowerPoint files (.pptx) — one document per slide."""

    @property
    def supported_types(self) -> list[str]:
        return [".pptx"]

    def load(self, source: str) -> list[LoadedDocument]:
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("pip install python-pptx")

        prs = Presentation(source)
        docs = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            content = "\n".join(texts)
            if content:
                docs.append(LoadedDocument(
                    content=content,
                    metadata={"slide": i, "source": source},
                    source=source,
                    provider="langchain",
                    loader="pptx",
                ))
        return docs
